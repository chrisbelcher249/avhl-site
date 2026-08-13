#!/usr/bin/env python3
from __future__ import annotations

import csv
import json
import re
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
TEAM_SOURCE = ROOT / 'data' / 'source' / 'major-team-specifications-2026-27.csv'
PLAYER_SOURCE = ROOT / 'data' / 'source' / 'returning-players-2026-27.xlsx'
TEAM_OUTPUT = ROOT / 'data' / 'teams.js'
PLAYER_OUTPUT = ROOT / 'data' / 'returningPlayers.js'

DIVISIONS = {
    'Pacific': [
        'Arizona Heat', 'Denver Mountain Lions', 'Honolulu Hawks', 'Las Vegas Vipers',
        'Montana Miners', 'Nevada Archers', 'North Dakota Bison', 'Portland Sea Lions',
        'Seattle Dragons', 'Washington Admirals',
    ],
    'Central': [
        'Fort Worth Defenders', 'Houston Hammerheads', 'Iowa Rebels', 'Kansas City Metrostars',
        'Lincoln Lumberjacks', 'Memphis Cannons', 'Oklahoma Twisters', 'San Antonio Bandits',
        'South Dakota Spartans', 'St. Louis Leopards',
    ],
    'Atlantic': [
        'Atlanta Cobalts', 'Chicago Rockets', 'Cincinnati Thunderbolts', 'Cleveland Demons',
        'Columbus Cougars', 'Detroit Motors', 'Indianapolis Ghosts', 'New Orleans Whalers',
        'Puerto Rico Toros', 'Tennessee Wolverines',
    ],
    'Metropolitan': [
        'Baltimore Oceanics', 'Brooklyn Bulldogs', 'Charleston Tsunami', 'Florida Sunshine',
        'Jacksonville Blood Hounds', 'Long Island Blizzard', 'New York Steamrollers',
        'Philadelphia Destroyers', 'Raleigh Wildcats', 'Richmond Robbers',
    ],
}
DIVISION_ORDER = ['Pacific', 'Central', 'Atlantic', 'Metropolitan']
TEAM_DIVISION = {team: division for division, names in DIVISIONS.items() for team in names}


def clean(value):
    if value is None:
        return None
    value = str(value).strip()
    return None if value == '' or value.lower() == 'nan' else value


def slugify(value: str) -> str:
    return re.sub(r'[^a-z0-9]+', '-', value.lower()).strip('-')


def rgb_to_hex(value):
    value = clean(value)
    if not value:
        return None
    if value.startswith('#') and len(value) == 7:
        return value.upper()
    parts = re.findall(r'\d+', value)
    if len(parts) >= 3:
        r, g, b = [max(0, min(255, int(part))) for part in parts[:3]]
        return f'#{r:02X}{g:02X}{b:02X}'
    return None


def js_dump(value):
    return json.dumps(value, indent=2, ensure_ascii=False)


# ---- Team data ----
with TEAM_SOURCE.open(newline='', encoding='utf-8-sig') as handle:
    rows = list(csv.DictReader(handle))

teams = []
for row in rows:
    city = clean(row.get('City/Location'))
    nickname = clean(row.get('Team Nickname'))
    name = f'{city} {nickname}'
    if name not in TEAM_DIVISION:
        raise ValueError(f'No division mapping for {name!r}')
    division = TEAM_DIVISION[name]
    abbreviation = clean(row.get('Abbreviation'))
    teams.append({
        'name': name,
        'slug': slugify(name),
        'city': city,
        'nickname': nickname,
        'abbreviation': abbreviation,
        'playByPlayName': clean(row.get('Play by Play Team Name')),
        'conference': 'Western' if division in {'Pacific', 'Central'} else 'Eastern',
        'division': division,
        'arena': clean(row.get('Arena Name')),
        'mascot': {
            'name': clean(row.get('Mascot Name')),
            'number': clean(row.get('Mascot #')),
        },
        'presentation': {
            'goalHorn': clean(row.get('Goal Horn')),
            'goalSong': clean(row.get('Goal Song')),
            'winSong': clean(row.get('Win Song')),
            'winPresentation': clean(row.get('Win Presentation')),
        },
        'colors': {
            'primary': clean(row.get('Hex 1')) or rgb_to_hex(row.get('Team Color 1')) or '#000B36',
            'primaryName': clean(row.get('Color 1')),
            'secondary': rgb_to_hex(row.get('Team Color 2')) or '#FFFFFF',
            'secondaryName': clean(row.get('Color 2')),
            'tertiary': rgb_to_hex(row.get('Team Color 3')) or '#A90117',
            'tertiaryName': clean(row.get('Color 3')),
        },
        'assets': {
            'logo': f'/teams/{abbreviation}/logo.webp',
            'arena': f'/teams/{abbreviation}/arena.webp',
            'mascot': f'/teams/{abbreviation}/mascot.webp',
            'home': f'/teams/{abbreviation}/home.webp',
            'away': f'/teams/{abbreviation}/away.webp',
            'alt': f'/teams/{abbreviation}/alt.webp',
        },
    })

order = {name: idx for idx, name in enumerate(DIVISION_ORDER)}
teams.sort(key=lambda t: (order[t['division']], t['name']))
assert len(teams) == 40
assert len({t['abbreviation'] for t in teams}) == 40
assert all(sum(1 for t in teams if t['division'] == d) == 10 for d in DIVISION_ORDER)

team_js = f'''// AUTO-GENERATED from data/source/major-team-specifications-2026-27.csv
// Run: python scripts/generate-site-data.py

export const divisionOrder = {js_dump(DIVISION_ORDER)};

export const conferenceMap = {{
  Western: ["Pacific", "Central"],
  Eastern: ["Atlantic", "Metropolitan"],
}};

export const teams = {js_dump(teams)};

export const teamBySlug = Object.fromEntries(teams.map((team) => [team.slug, team]));
export const teamByAbbreviation = Object.fromEntries(teams.map((team) => [team.abbreviation, team]));
export const teamsByDivision = Object.fromEntries(
  divisionOrder.map((division) => [division, teams.filter((team) => team.division === division)])
);
'''
TEAM_OUTPUT.write_text(team_js, encoding='utf-8')

# ---- Returning player data ----
TEAM_NAMES = {t['name'] for t in teams}
wb = load_workbook(PLAYER_SOURCE, read_only=False, data_only=True)
if set(wb.sheetnames) != TEAM_NAMES:
    missing = sorted(TEAM_NAMES - set(wb.sheetnames))
    extra = sorted(set(wb.sheetnames) - TEAM_NAMES)
    raise ValueError(f'Workbook/team mismatch. Missing={missing}, extra={extra}')


def header_map(ws, row_number):
    return {clean(cell.value): idx for idx, cell in enumerate(ws[row_number], start=1) if clean(cell.value)}


def norm(value):
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def value_at(ws, row_number, mapping, label):
    col = mapping.get(label)
    return norm(ws.cell(row_number, col).value) if col else None


def as_id(value):
    if value is None:
        return None
    text = str(value).strip()
    return text.zfill(4) if text.isdigit() else text


def make_skater(ws, r, h):
    return {
        'id': as_id(value_at(ws, r, h, 'AVHL ID')),
        'name': clean(value_at(ws, r, h, 'Full Name')),
        'yearsLeft': value_at(ws, r, h, 'Years Left'),
        'season': {
            'g': value_at(ws, r, h, '25-26 G') or 0,
            'a': value_at(ws, r, h, '25-26 A') or 0,
            'pts': value_at(ws, r, h, '25-26 P') or 0,
        },
        'career': {
            'g': value_at(ws, r, h, 'Total G') or 0,
            'a': value_at(ws, r, h, 'Total A') or 0,
            'pts': value_at(ws, r, h, 'Total PTS') or 0,
        },
        'overall': value_at(ws, r, h, 'Overall Rating'),
        'birthdate': value_at(ws, r, h, 'Birthdate'),
        'number': value_at(ws, r, h, 'Jersey #'),
        'position': clean(value_at(ws, r, h, 'Position(s)')),
        'height': clean(value_at(ws, r, h, 'Height')),
        'weight': value_at(ws, r, h, 'Weight'),
        'age': value_at(ws, r, h, 'Age'),
        'shot': clean(value_at(ws, r, h, 'Shot')),
        'playerType': clean(value_at(ws, r, h, 'Player Type')),
        'nhlTeam': clean(value_at(ws, r, h, 'Team')),
        'ratings': {
            'Deking': value_at(ws, r, h, 'Deking'),
            'Passing': value_at(ws, r, h, 'Passing'),
            'Puck Control': value_at(ws, r, h, 'Puck Control'),
            'Off. Awareness': value_at(ws, r, h, 'Off. Awareness'),
            'Wrist Shot Acc.': value_at(ws, r, h, 'Wrist Shot Accuracy'),
            'Def. Awareness': value_at(ws, r, h, 'Def. Awareness'),
            'Faceoffs': value_at(ws, r, h, 'Faceoffs'),
            'Stick Checking': value_at(ws, r, h, 'Stick Checking'),
            'Acceleration': value_at(ws, r, h, 'Acceleration'),
            'Speed': value_at(ws, r, h, 'Speed'),
            'Body Checking': value_at(ws, r, h, 'Body Checking'),
            'Strength': value_at(ws, r, h, 'Strength'),
        },
    }


def make_goalie(ws, r, h):
    def pct(label):
        v = value_at(ws, r, h, label)
        return round(float(v), 3) if v not in (None, '') else None
    return {
        'id': as_id(value_at(ws, r, h, 'AVHL ID')),
        'name': clean(value_at(ws, r, h, 'Full Name')),
        'yearsLeft': value_at(ws, r, h, 'Years Left'),
        'season': {
            'sa': value_at(ws, r, h, '25-26 SA') or 0,
            'sv': value_at(ws, r, h, '25-26 SV') or 0,
            'svPct': pct('25-26 SV%'),
        },
        'career': {
            'sa': value_at(ws, r, h, 'Total SA') or 0,
            'sv': value_at(ws, r, h, 'Total SV') or 0,
            'svPct': pct('Career SV%'),
        },
        'overall': value_at(ws, r, h, 'Overall Rating'),
        'birthdate': value_at(ws, r, h, 'Birthdate'),
        'number': value_at(ws, r, h, 'Jersey #'),
        'position': 'G',
        'height': clean(value_at(ws, r, h, 'Height')),
        'weight': value_at(ws, r, h, 'Weight'),
        'age': value_at(ws, r, h, 'Age'),
        'glove': clean(value_at(ws, r, h, 'Glove')),
        'nhlTeam': clean(value_at(ws, r, h, 'Team')),
        'ratings': {
            'Angles': value_at(ws, r, h, 'Angles'),
            'Breakaway': value_at(ws, r, h, 'Breakaway'),
            'Five Hole': value_at(ws, r, h, 'Five Hole'),
            'Glove High': value_at(ws, r, h, 'Glove High'),
            'Glove Low': value_at(ws, r, h, 'Glove Low'),
            'Stick High': value_at(ws, r, h, 'Stick High'),
            'Stick Low': value_at(ws, r, h, 'Stick Low'),
            'Rebound Control': value_at(ws, r, h, 'Rebound Control'),
            'Recover': value_at(ws, r, h, 'Recover'),
            'Agility': value_at(ws, r, h, 'Agility'),
            'Speed': value_at(ws, r, h, 'Speed'),
            'Vision': value_at(ws, r, h, 'Vision'),
        },
    }

returning = {}
total_players = 0
for ws in wb.worksheets:
    skater_header = goalie_header = None
    for row in range(1, ws.max_row + 1):
        first = clean(ws.cell(row, 1).value)
        if first == 'AVHL ID':
            # First such row belongs to skaters, second to goalies.
            if skater_header is None:
                skater_header = row
            else:
                goalie_header = row
    skaters, goalies = [], []
    if skater_header:
        h = header_map(ws, skater_header)
        r = skater_header + 1
        while r <= ws.max_row:
            first = clean(ws.cell(r, 1).value)
            if not first or str(first).startswith('RETURNING GOALIES'):
                break
            if str(first).isdigit():
                skaters.append(make_skater(ws, r, h))
            r += 1
    if goalie_header:
        h = header_map(ws, goalie_header)
        r = goalie_header + 1
        while r <= ws.max_row:
            first = clean(ws.cell(r, 1).value)
            if not first:
                break
            if str(first).isdigit():
                goalies.append(make_goalie(ws, r, h))
            r += 1
    total_players += len(skaters) + len(goalies)
    returning[slugify(ws.title)] = {
        'teamName': ws.title,
        'count': len(skaters) + len(goalies),
        'skaters': skaters,
        'goalies': goalies,
    }

assert len(returning) == 40
assert total_players == 555, f'Expected 555 returning players, found {total_players}'

player_js = f'''// AUTO-GENERATED from data/source/returning-players-2026-27.xlsx
// Run: python scripts/generate-site-data.py

export const returningPlayersBySlug = {js_dump(returning)};

export const returningPlayerCount = Object.values(returningPlayersBySlug)
  .reduce((sum, roster) => sum + roster.count, 0);
'''
PLAYER_OUTPUT.write_text(player_js, encoding='utf-8')
print(f'Generated {TEAM_OUTPUT.relative_to(ROOT)} for {len(teams)} teams')
print(f'Generated {PLAYER_OUTPUT.relative_to(ROOT)} for {total_players} returning players')
